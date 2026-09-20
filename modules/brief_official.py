# -*- coding: utf-8 -*-
"""按厦门海洋预报台**官方简报版式**生成 Word 简报。

版式参数取自课题存档的官方样件（厦门中心简报材料/）：
    2026044厦门海浪Ⅳ级蓝色警报.docx、2026030漳州核电海浪Ⅳ级蓝色警报.docx、
    2026002第三东通项目道海浪消息.docx、2026_09_02漳州核电海域海洋环境预报.pdf、
    风暴潮、海浪形势分析预测（2021年第014期）.pdf

复刻要点：
    · 页面 A4，页边距 上 2.0 / 下 2.0 / 左 2.2 / 右 2.2 cm
    · 抬头（发文单位）黑体 18pt 加粗，三行
    · 大标题 黑体 36pt 加粗居中（海浪警报 / 海 浪 消 息 / 风暴潮警报 …）
    · 时间 12pt：时间：2026年09月18日14时
    · 编号 + 签发 12pt 同一行：编号：海浪2026-045　　　　签发：张世民
    · 依据段 14pt 加粗：…根据《…应急预案（风暴潮、海浪、海啸）》发布…。
    · 小标题 黑体 16pt 居中
    · 正文 14pt
    · 表格：风暴潮/形势分析 → 潮位×警戒潮位表；海洋环境预报 → 风浪时效表
    · 预警提示 宋体 12pt

编号自动递增：记录在 <briefs>/.serial.json（按 类别+年份 计数）。
"""
from __future__ import annotations

import datetime
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HEI = "黑体"
SONG = "宋体"

# 三行抬头（可由 data["agencies"] 覆盖）
AGENCIES = ["厦门市海洋发展局", "自然资源部厦门海洋预报台", "厦门海洋环境预报台"]
ISSUER = "张世民"                      # 签发人（可由 data["issuer"] 覆盖）
BASIS_NAME = "《自然资源部厦门海洋中心海洋灾害应急执行预案（风暴潮、海浪、海啸）》"

# 大标题（官方按类型固定）
TITLES = {
    "storm_surge_alert": "风暴潮警报",
    "wave_alert": "海浪警报",
    "wave_message": "海 浪 消 息",
    "marine_env_forecast": "海洋环境预报",
    "situation_analysis": "风暴潮、海浪形势分析预测",
}
# 编号前缀
SERIAL_PREFIX = {
    "storm_surge_alert": "风暴潮",
    "wave_alert": "海浪",
    "wave_message": "海浪",
    "marine_env_forecast": "海洋环境预报",
    "situation_analysis": "形势分析",
}


def _clean(text: Any) -> str:
    """去掉 Markdown 标记（官方简报是纯文本，不能出现 ** 之类）。"""
    s = str(text or "")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"^#+\s*", "", s, flags=re.M)
    s = re.sub(r"^[⏰📄⚠️✅❌]\s*", "", s, flags=re.M)
    return s.strip()


def _font(run, name: str, size: float, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    return run


def _para(doc, text: str, font: str = SONG, size: float = 14, bold: bool = False,
          align=None, indent_pt: float = 0, spacing: float = 1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    if indent_pt:
        pf.first_line_indent = Pt(indent_pt)
    _font(p.add_run(_clean(text)), font, size, bold)
    return p


def next_serial(prefix: str, year: int, briefs_dir: Path) -> str:
    """生成编号：<前缀><年>-<3位序号>，序号按 .serial.json 递增。"""
    f = briefs_dir / ".serial.json"
    data: Dict[str, int] = {}
    if f.exists():
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    key = f"{prefix}{year}"
    data[key] = int(data.get(key, 0)) + 1
    try:
        briefs_dir.mkdir(parents=True, exist_ok=True)
        f.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass
    return f"{prefix}{year}-{data[key]:03d}"


def render(data: Dict[str, Any], output_dir, now: Optional[datetime.datetime] = None) -> Path:
    """生成官方版式 Word，返回文件路径。"""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    now = now or datetime.datetime.now()

    tpl = str(data.get("template_type") or "storm_surge_alert")
    big_title = str(data.get("big_title") or TITLES.get(tpl, "海洋预警简报"))
    sub_title = _clean(data.get("title") or big_title)      # 小标题（含区域/级别）
    agencies = data.get("agencies") or AGENCIES
    issuer = str(data.get("issuer") or ISSUER)
    serial = str(data.get("serial") or
                 next_serial(SERIAL_PREFIX.get(tpl, "简报"), now.year, output_dir))

    # 文件名：编号 + 小标题，便于归档
    safe = "".join(ch for ch in f"{serial}_{sub_title}" if ch not in r'\/:*?"<>|')
    out = output_dir / f"{safe}.docx"

    doc = Document()

    # 页面：A4 + 官方边距
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.2)

    st = doc.styles["Normal"]
    st.font.name = SONG
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    # ① 抬头（黑体 18pt 加粗，可多行）
    for a in (agencies if isinstance(agencies, (list, tuple)) else [agencies]):
        _para(doc, a, HEI, 18, True, spacing=1.2)

    # ② 大标题（黑体 36pt 居中）
    _para(doc, big_title, HEI, 36, True, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.2)

    # ③ 时间（12pt）
    _para(doc, f"时间：{now:%Y年%m月%d日%H时}", SONG, 12, spacing=1.2)

    # ④ 编号 + 签发（12pt 同行）
    _para(doc, f"编号：{serial}　　　　　　　　　　　　签发：{issuer}", SONG, 12, spacing=1.2)

    # ⑤ 依据段（14pt 加粗）
    basis = _clean(data.get("basis") or "")
    if basis and not basis.startswith("《"):
        basis_text = basis
    else:
        name = basis or BASIS_NAME
        publisher = "厦门海洋环境预报台" if len(agencies) > 2 else "自然资源部厦门海洋预报台"
        basis_text = f"{publisher}根据{name}发布{sub_title}。"
    _para(doc, basis_text, SONG, 14, True)

    # ⑥ 小标题（黑体 16pt 居中）
    _para(doc, sub_title, HEI, 16, False, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ⑦ 正文（14pt，可多段）
    body = data.get("summary") or data.get("body") or ""
    if isinstance(body, (list, tuple)):
        for b in body:
            _para(doc, b, SONG, 14, indent_pt=28)
    elif body:
        for line in str(body).split("\n"):
            if line.strip():
                _para(doc, line, SONG, 14, indent_pt=28)

    # ⑧ 表格（按类型）
    if tpl in ("storm_surge_alert", "situation_analysis"):
        _table_stations(doc, data)
    elif tpl == "marine_env_forecast":
        _table_marine(doc, data)
    elif tpl in ("wave_alert", "wave_message") and len(data.get("segments") or []) >= 2:
        _table_segments(doc, data)

    # ⑨ 关注提示 + 预警提示
    if data.get("notice"):
        _para(doc, data["notice"], SONG, 14, indent_pt=28)
    if data.get("tip"):
        _para(doc, data["tip"], SONG, 12)

    # ⑩ 备注 / 落款（可选）
    if data.get("note"):
        _para(doc, data["note"], SONG, 12)
    for line in (data.get("footer") or []):
        if line:
            _para(doc, line, SONG, 12, spacing=1.2)

    doc.save(str(out))
    return out


def _table_stations(doc, data: Dict[str, Any]) -> None:
    rows = data.get("stations") or []
    if not rows:
        return
    headers = ["站位", "日期", "时间", "高潮位(cm)", "警戒潮位(cm)", "预警级别"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _font(cell.paragraphs[0].add_run(h), SONG, 12, True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in rows:
        cells = t.add_row().cells
        vals = [s.get("station", ""), s.get("date", ""), s.get("time", ""),
                str(s.get("high_tide_cm", "")), str(s.get("warn", "")),
                str(s.get("level", ""))]
        for i, v in enumerate(vals):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(_clean(v)), SONG, 12)


def _table_segments(doc, data: Dict[str, Any]) -> None:
    rows = data.get("segments") or []
    headers = ["海域", "浪高(m)", "预警级别"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), SONG, 12, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in rows:
        cells = t.add_row().cells
        vals = [s.get("area", ""), str(s.get("wave", "")), str(s.get("level", ""))]
        for i, v in enumerate(vals):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(_clean(v)), SONG, 12)


def _table_marine(doc, data: Dict[str, Any]) -> None:
    """海洋环境预报：时效 × 要素（风/浪）表。数据缺失则不生成。"""
    rows = data.get("forecast_table") or []
    if not rows:
        return
    headers = ["时效", "要素", "24 小时", "48 小时", "72 小时"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), SONG, 12, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        cells = t.add_row().cells
        vals = [r.get("label", ""), r.get("element", ""), r.get("h24", ""),
                r.get("h48", ""), r.get("h72", "")]
        for i, v in enumerate(vals):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(_clean(v)), SONG, 12)
