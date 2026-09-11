#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Markdown → Word（中文公文版式），用于把情况说明等材料转成可直接提交的 .docx。

用法：
    python scripts/md_to_docx.py docs\情况说明-DeepSeek公共API申请.md
    python scripts/md_to_docx.py 输入.md -o 输出.docx --title "关于……的情况说明"

支持的 Markdown 语法（覆盖本项目文档用到的部分）：
    # / ## / ### 标题、正文段落、- 无序列表、1. 有序列表、
    | 表格 |、**加粗**、> 引用、--- 分隔线、`行内代码`
版式：标题黑体、正文仿宋_GB2312 小三、行距 1.5、首行缩进 2 字符，A4 页边距。
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("需要 python-docx：pip install python-docx")

CN_BODY = "仿宋_GB2312"
CN_HEAD = "黑体"
CN_TITLE = "方正小标宋简体"


def _set_font(run, name: str, size: float, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _style_doc(doc: Document):
    """A4 + 公文页边距 + 默认正文样式。"""
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(3.7), Cm(3.5)
        s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)
    st = doc.styles["Normal"]
    st.font.name = CN_BODY
    st.font.size = Pt(15)               # 小三
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)


def _add_runs(par, text: str, font: str, size: float, base_bold: bool = False):
    """处理 **加粗** 与 `行内代码`。"""
    for seg in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            _set_font(par.add_run(seg[2:-2]), font, size, True)
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1])
            _set_font(r, "Consolas", size - 1)
            r.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        else:
            _set_font(par.add_run(seg), font, size, base_bold)


def _split_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def convert(md_path: Path, out_path: Path, title: str = "") -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _style_doc(doc)

    i = 0
    first_h1_done = False
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # ---- 表格 ----
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = _split_row(s)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                _add_runs(cell.paragraphs[0], h, CN_HEAD, 10.5, True)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for row in rows:
                cells = t.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    cells[j].text = ""
                    _add_runs(cells[j].paragraphs[0], val, CN_BODY, 10.5)
            doc.add_paragraph()
            continue

        # ---- 标题 ----
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            p = doc.add_paragraph()
            if level == 1 and not first_h1_done:
                first_h1_done = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_font(p.add_run(title or text), CN_TITLE, 22, False)
                p.paragraph_format.space_after = Pt(18)
            elif level == 1:
                _set_font(p.add_run(text), CN_HEAD, 16, False)
                p.paragraph_format.space_before = Pt(12)
            elif level == 2:
                _set_font(p.add_run(text), CN_HEAD, 15, False)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
            else:
                _set_font(p.add_run(text), CN_HEAD, 14, False)
                p.paragraph_format.space_before = Pt(6)
            i += 1
            continue

        # ---- 分隔线 ----
        if re.match(r"^-{3,}$", s):
            i += 1
            continue

        # ---- 引用 ----
        if s.startswith(">"):
            body = s.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _add_runs(p, body, CN_BODY, 10.5)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ---- 列表 ----
        m = re.match(r"^([-*])\s+(.*)$", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            _add_runs(p, "● " + m.group(2), CN_BODY, 12)
            i += 1
            continue
        m = re.match(r"^(\d+)[.、]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            _add_runs(p, f"{m.group(1)}. {m.group(2)}", CN_BODY, 12)
            i += 1
            continue

        # ---- 空行 ----
        if not s:
            i += 1
            continue

        # ---- 正文段落 ----
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(30)          # 首行缩进 2 字符
        pf.line_spacing = 1.5
        pf.space_after = Pt(4)
        _add_runs(p, s, CN_BODY, 15)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Markdown → Word（中文公文版式）")
    ap.add_argument("src", help="输入的 .md 文件")
    ap.add_argument("-o", "--out", default="", help="输出的 .docx（缺省与输入同目录同名）")
    ap.add_argument("--title", default="", help="覆盖一级标题（用于正式文件名与标题不一致的情况）")
    args = ap.parse_args()

    src = Path(args.src)
    if not src.exists():
        sys.exit(f"找不到文件：{src}")
    out = Path(args.out) if args.out else src.with_suffix(".docx")
    p = convert(src, out, args.title)
    print(f"已生成：{p}")


if __name__ == "__main__":
    main()
